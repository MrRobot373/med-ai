"""Text extraction per file type. Each parser returns a list of pages:
`{"page": int, "text": str}` (page numbers are 1-based; non-paginated formats are one
page per sheet/section). Images return no text here — they go through vision."""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass

SUPPORTED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "text",
    "text/markdown": "text",
    "text/csv": "csv",
    "application/json": "json",
    "image/png": "image",
    "image/jpeg": "image",
    "image/webp": "image",
    "application/dicom": "dicom",
}
EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".dcm": "application/dicom",
    ".dicom": "application/dicom",
}
MAGIC = {
    "pdf": [b"%PDF"],
    "docx": [b"PK"],
    "xlsx": [b"PK"],
    "image": [b"\x89PNG", b"\xff\xd8\xff", b"RIFF"],
}


@dataclass
class Parsed:
    pages: list[dict]
    kind: str  # parser family

    @property
    def text(self) -> str:
        return "\n\n".join(p["text"] for p in self.pages if p["text"].strip())


def detect_mime(filename: str, declared: str | None) -> str | None:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in EXTENSIONS:
        return EXTENSIONS[ext]
    if declared in SUPPORTED:
        return declared
    return None


def verify_magic(family: str, data: bytes) -> bool:
    if family == "dicom":
        from app.ingestion.dicom import is_dicom

        return is_dicom(data)
    sigs = MAGIC.get(family)
    return True if not sigs else any(data.startswith(s) for s in sigs)


def parse(data: bytes, mime: str) -> Parsed:
    family = SUPPORTED[mime]
    if family == "pdf":
        return Parsed(_pdf(data), family)
    if family == "docx":
        return Parsed(_docx(data), family)
    if family == "xlsx":
        return Parsed(_xlsx(data), family)
    if family == "csv":
        return Parsed([{"page": 1, "text": _csv(data)}], family)
    if family == "json":
        return Parsed([{"page": 1, "text": _json(data)}], family)
    if family in ("image", "dicom"):
        return Parsed([], family)
    return Parsed([{"page": 1, "text": data.decode("utf-8", errors="replace")}], family)


def _pdf(data: bytes) -> list[dict]:
    import pymupdf

    pages = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            text = _layout_text(page.get_text("words")) or page.get_text("text") or ""
            pages.append({"page": i, "text": text})
    return pages


def _layout_text(words: list[tuple]) -> str:
    """Rebuilds printed lines from word positions. Words on one baseline form a line;
    a wide horizontal gap becomes " | " so column layouts (lab tables) keep their cells:
    "Haemoglobin | 12.9 | g/dL | 13.0 - 17.0 | L". Prose keeps ordinary spacing."""
    if not words:
        return ""
    rows: list[list[tuple]] = []
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        mid = (w[1] + w[3]) / 2
        if rows and abs(mid - rows[-1][0][4]) <= max(2.5, (w[3] - w[1]) * 0.35):
            rows[-1][1].append(w)
        else:
            rows.append([(0, 0, 0, 0, mid), [w]])
    lines = []
    for _anchor, row in rows:
        row.sort(key=lambda w: w[0])
        height = max(w[3] - w[1] for w in row) or 10
        parts = [row[0][4]]
        for prev, cur in zip(row, row[1:], strict=False):
            gap = cur[0] - prev[2]
            sep = " | " if gap > height * 1.6 and prev[4] != "|" and cur[4] != "|" else " "
            parts.append(sep + cur[4])
        lines.append("".join(parts))
    return "\n".join(lines)


def _docx(data: bytes) -> list[dict]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower() if para.style is not None else ""
            parts.append(("## " if style.startswith("heading") else "") + para.text.strip())
    for table in d.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        parts.append("\n".join(rows))
    return [{"page": 1, "text": "\n".join(parts)}]


def _xlsx(data: bytes) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        pages.append({"page": i, "text": f"## {ws.title}\n" + "\n".join(rows)})
    return pages


def _csv(data: bytes) -> str:
    reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
    return "\n".join(" | ".join(r) for r in reader)


_TITLE_KEYS = ("title", "name", "disease", "condition", "drug", "topic", "heading")


def _render_value(value, indent: int = 0) -> str:
    pad = "  " * indent
    if isinstance(value, list):
        return "\n".join(f"{pad}- {_render_value(v, indent + 1).strip()}" for v in value)
    if isinstance(value, dict):
        return "\n".join(
            f"{pad}{k.replace('_', ' ').capitalize()}: {_render_value(v, indent + 1).strip()}" for k, v in value.items()
        )
    return f"{pad}{value}"


def _json(data: bytes) -> str:
    """Structured knowledge packs (a list of records) become one readable section per
    record so each chunk stays a coherent topic."""
    try:
        value = json.loads(data)
    except ValueError:
        return data.decode("utf-8", errors="replace")
    if isinstance(value, dict):
        for key in ("items", "entries", "records", "data"):
            if isinstance(value.get(key), list):
                value = value[key]
                break
    if isinstance(value, list) and value and all(isinstance(v, dict) for v in value):
        sections = []
        for record in value:
            title = next((str(record[k]) for k in _TITLE_KEYS if record.get(k)), "Entry")
            body = "\n".join(
                f"{k.replace('_', ' ').capitalize()}: {_render_value(v, 1).strip()}"
                for k, v in record.items()
                if str(record.get(k)) != title
            )
            sections.append(f"## {title}\n{body}")
        return "\n\n".join(sections)
    return json.dumps(value, indent=2, ensure_ascii=False)
