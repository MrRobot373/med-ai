"""DOCX text extraction.

DOCX has no reliable page-boundary concept without a full layout engine — spec §51
explicitly allows falling back to section/paragraph provenance instead of a page
number for formats like this, so the whole document is returned as one logical unit
rather than a fabricated page count.
"""

from __future__ import annotations

from io import BytesIO

import docx


def extract_docx_text(content: bytes) -> str:
    document = docx.Document(BytesIO(content))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        rows = [r for r in rows if r.strip(" |")]
        if rows:
            parts.append("[Table]\n" + "\n".join(rows))

    return "\n\n".join(parts)
