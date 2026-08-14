"""Synthetic document generators — used by pytest AND scripts/generate_sample_reports.py
(which imports this module directly, see that script for how). Kept in one place so
the "what does a realistic test PDF/DOCX/XLSX/image look like" logic isn't duplicated.

Nothing here is real patient data — every value is fabricated for testing/demo purposes.
"""

from __future__ import annotations

from io import BytesIO

import docx
import numpy as np
import openpyxl
import pymupdf
from PIL import Image, ImageDraw
from pydicom.dataset import FileDataset, FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian, generate_uid


def make_pdf_bytes(lines: list[str]) -> bytes:
    """A normal, real-text-layer PDF — the common case."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 50
    for line in lines:
        page.insert_text((50, y), line, fontsize=11)
        y += 16
    data = doc.tobytes()
    doc.close()
    return data


def make_scanned_pdf_bytes(lines: list[str]) -> bytes:
    """A PDF whose only content is a rendered image (no text layer) — simulates a
    scanned document to exercise the OCR-fallback path."""
    image_bytes = make_png_bytes(lines, size=(800, 600))

    doc = pymupdf.open()
    page = doc.new_page(width=800, height=600)
    page.insert_image(pymupdf.Rect(0, 0, 800, 600), stream=image_bytes)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes(sheet_name: str, rows: list[list]) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_dicom_bytes(modality: str = "OP", rows: int = 8, columns: int = 8) -> bytes:
    """A minimal-but-valid synthetic DICOM file (random noise pixel data) — enough to
    exercise detection/metadata-reading/preview-rendering without a real DICOM sample."""
    file_meta = FileMetaDataset()
    file_meta.MediaStorageSOPClassUID = generate_uid()
    file_meta.MediaStorageSOPInstanceUID = generate_uid()
    file_meta.TransferSyntaxUID = ExplicitVRLittleEndian

    dataset = FileDataset(None, {}, file_meta=file_meta, preamble=b"\x00" * 128)
    dataset.PatientName = "Test^Patient"
    dataset.PatientID = "DEMO-0000"
    dataset.Modality = modality
    dataset.StudyDescription = "Test Study"
    dataset.SeriesDescription = "Test Series"
    dataset.Rows = rows
    dataset.Columns = columns
    dataset.SamplesPerPixel = 1
    dataset.PhotometricInterpretation = "MONOCHROME2"
    dataset.BitsAllocated = 8
    dataset.BitsStored = 8
    dataset.HighBit = 7
    dataset.PixelRepresentation = 0
    dataset.PixelData = (np.random.rand(rows, columns) * 255).astype("uint8").tobytes()

    buffer = BytesIO()
    # little_endian/implicit_vr are inferred from file_meta.TransferSyntaxUID (set
    # above) — setting the deprecated is_little_endian/is_implicit_VR attributes
    # directly is unnecessary here and pydicom 3.x warns on them.
    dataset.save_as(buffer, enforce_file_format=True)
    return buffer.getvalue()


def make_png_bytes(lines: list[str], size: tuple[int, int] = (800, 400)) -> bytes:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    y = 20
    for line in lines:
        draw.text((20, y), line, fill="black")
        y += 24
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
